from __future__ import annotations

from typing import Dict
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from domain.bom_models import TotalEntry


def export_totals_to_excel(totals: Dict[str, TotalEntry], filepath: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Totalizzazione"

    headers = [
        "Codice Interno",
        "Descrizione",
        "Quantità",
        "UM",
        "Manufacturer",
        "Codice Produttore",
    ]
    ws.append(headers)

    for _, entry in totals.items():
        ws.append([
            entry.internal_code,
            entry.description,
            entry.qty,
            entry.unit,
            entry.manufacturer,
            entry.manufacturer_code,
        ])

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    wb.save(filepath)
